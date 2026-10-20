from pathlib import Path
import re
import os

import faiss
import numpy as np
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer
from pptx import Presentation

from rag.topic_profiles import TOPIC_PROFILES


class Retriever:
    """Small BGE + FAISS index built from files in the knowledge folder."""

    def __init__(self, folder):
        self.folder = Path(folder)
        self.folder.mkdir(exist_ok=True)
        self.chunks = []
        self.query_cache = {}
        self.offline = os.getenv("HINTOS_OFFLINE") == "1"
        model_folder = Path(__file__).parent.parent / "models" / "bge-small-en-v1.5"
        if model_folder.exists():
            self.model = SentenceTransformer(str(model_folder), local_files_only=self.offline)
        elif self.offline:
            raise RuntimeError(
                "Offline mode is on, but the BGE model is missing. Run prepare_model.py before going offline."
            )
        else:
            self.model = SentenceTransformer("BAAI/bge-small-en-v1.5")
        self._load_documents()
        self.topic_names = list(TOPIC_PROFILES)
        self.topic_vectors = np.asarray(
            self.model.encode(
                [
                    "Represent this sentence for searching relevant passages: " + TOPIC_PROFILES[name]
                    for name in self.topic_names
                ],
                normalize_embeddings=True,
            ),
            dtype="float32",
        )
        self.index = None
        if self.chunks:
            vectors = self.model.encode(
                ["Represent this sentence for searching relevant passages: " + c["text"] for c in self.chunks],
                normalize_embeddings=True,
            )
            vectors = np.asarray(vectors, dtype="float32")
            self.index = faiss.IndexFlatIP(vectors.shape[1])
            self.index.add(vectors)

    def _read_file(self, path):
        suffix = path.suffix.lower()
        if suffix in [".txt", ".md"]:
            return path.read_text(encoding="utf-8", errors="ignore")
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        if suffix == ".pptx":
            presentation = Presentation(str(path))
            return "\n".join(
                shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text")
            )
        return ""

    def _load_documents(self):
        for path in self.folder.iterdir():
            if path.is_file():
                text = re.sub(r"\s+", " ", self._read_file(path)).strip()
                words = text.split()
                # Small overlapping chunks keep a little context at the edges.
                for start in range(0, len(words), 100):
                    piece = " ".join(words[start:start + 120])
                    if piece:
                        self.chunks.append({"text": piece, "source": path.name})

    def search(self, question, count=3):
        if self.index is None or not question.strip():
            return []
        vector = self._query_vector(question)
        scores, ids = self.index.search(vector, min(count, len(self.chunks)))
        results = []
        for score, chunk_id in zip(scores[0], ids[0]):
            if chunk_id >= 0:
                results.append({**self.chunks[int(chunk_id)], "score": float(score)})
        return results

    def classify_topic(self, question, minimum_score=0.54, minimum_margin=0.04):
        """Use the local embedding model to match meaning, not just shared words."""
        if not question.strip():
            return ""
        vector = self._query_vector(question)
        scores = vector @ self.topic_vectors.T
        ranked = np.argsort(scores[0])
        best = int(ranked[-1])
        second = int(ranked[-2])
        if float(scores[0][best]) < minimum_score or float(scores[0][best] - scores[0][second]) < minimum_margin:
            return ""
        return self.topic_names[best]

    def _query_vector(self, question):
        """Reuse the query embedding for topic matching and note retrieval."""
        vector = self.query_cache.get(question)
        if vector is None:
            encoded = self.model.encode(
                ["Represent this sentence for searching relevant passages: " + question],
                normalize_embeddings=True,
            )
            vector = np.asarray(encoded, dtype="float32")
            if len(self.query_cache) >= 64:
                self.query_cache.pop(next(iter(self.query_cache)))
            self.query_cache[question] = vector
        return vector
